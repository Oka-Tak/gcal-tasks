#!/usr/bin/env python3
"""Extract plain text from a binary document for the OneDrive → RAG sync.

Open WebUI's built-in loaders return empty for pptx/pdf/xlsx/csv on this box
(they route through unstructured/docling, which aren't installed) AND embedding
large docs on CPU saturates the machine. So owui-sync.ts extracts text here —
row-/char-capped so a fat spreadsheet can't blow up into millions of chars —
and pushes only the text to Open WebUI.

Usage: extract-text.py <path>
  stdout: extracted text (possibly empty)
  exit 0: extraction ran (empty stdout = genuinely no text → permanent skip)
  exit 2: extraction failed (deterministic, treated as permanent skip)

Run with the Open WebUI venv python (has pptx/pypdf/openpyxl).
"""
import os
import sys

MAX_CHARS = 200_000
MAX_ROWS = 2000


def emit(text: str) -> None:
    sys.stdout.write((text or "").strip()[:MAX_CHARS])


def extract_pptx(path: str) -> str:
    from pptx import Presentation

    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        chunk = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                chunk.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    chunk.append("\t".join(c.text for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            note = slide.notes_slide.notes_text_frame.text.strip()
            if note:
                chunk.append("[ノート] " + note)
        if chunk:
            parts.append("--- slide {} ---\n{}".format(i, "\n".join(chunk)))
    return "\n\n".join(parts)


def extract_pdf(path: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    parts = []
    total = 0
    for i, page in enumerate(reader.pages, 1):
        t = (page.extract_text() or "").strip()
        if t:
            parts.append("--- p{} ---\n{}".format(i, t))
            total += len(t)
        if total > MAX_CHARS:
            break
    return "\n\n".join(parts)


def extract_xlsx(path: str) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            cells = [str(c) for c in r if c is not None]
            if cells:
                rows.append("\t".join(cells))
            if len(rows) >= MAX_ROWS:
                break
        if rows:
            parts.append("# {}\n{}".format(ws.title, "\n".join(rows)))
    wb.close()
    return "\n\n".join(parts)


def extract_csv(path: str) -> str:
    import csv

    for enc in ("utf-8-sig", "cp932", "utf-8"):
        try:
            with open(path, encoding=enc, newline="") as f:
                rows = []
                for i, row in enumerate(csv.reader(f)):
                    rows.append("\t".join(row))
                    if i >= MAX_ROWS:
                        break
                return "\n".join(rows)
        except UnicodeDecodeError:
            continue
    return ""


EXTRACTORS = {
    ".pptx": extract_pptx,
    ".pdf": extract_pdf,
    ".xlsx": extract_xlsx,
    ".csv": extract_csv,
}


def main() -> int:
    if len(sys.argv) < 2:
        sys.stderr.write("usage: extract-text.py <path>")
        return 2
    path = sys.argv[1]
    ext = os.path.splitext(path)[1].lower()
    fn = EXTRACTORS.get(ext)
    if not fn:
        return 2
    try:
        emit(fn(path))
        return 0
    except Exception as e:  # deterministic failure → caller records permanent skip
        sys.stderr.write("{}: {}".format(type(e).__name__, e))
        return 2


if __name__ == "__main__":
    sys.exit(main())
